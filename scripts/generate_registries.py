from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

OUT = Path("03_Registries/Onboarding_Mapping_Agent_Knowledge_Registries.xlsx")
OUT.parent.mkdir(parents=True, exist_ok=True)

SPECS = [
    ("Taxonomy", [
        "Taxonomy ID","Status","Primary Group","Secondary Group","Canonical Mapping Name",
        "Definition","Defining Characteristics","Workfront Applicability","Supporting Evidence",
        "Confidence","Approved By","Approved Date","Supersedes","Notes"
    ], "Approved or candidate mapping-family taxonomy. Do not populate from filenames alone.", {
        "Status": ["Candidate","Approved","Superseded"],
        "Workfront Applicability": ["STANDARD","CONDITIONAL","EXCEPTION","NOT OBSERVED","UNKNOWN"],
        "Confidence": ["Strong","Moderate","Weak"]
    }),
    ("Decisions", [
        "Decision ID","Status","Decision Type","Mapping Family","Object / Area","Field",
        "Decision","Rationale","Evidence References","Requested / Raised By",
        "Approved By","Approved Date","Supersedes Decision ID","Notes"
    ], "Important approved/rejected analytical and business decisions that should influence future reasoning.", {
        "Status": ["Proposed","Approved","Rejected","Superseded"],
        "Decision Type": ["Taxonomy","Mapping","Workfront","Governance","Other"]
    }),
    ("Sources & Evidence", [
        "Source ID","Source Name","Source Type","Location","Historical Context","Partner / Reseller",
        "Concepts Found","Salesforce Objects","Workfront Relevance","Mapping Completeness",
        "Potential Mapping Family","Source Status","Confidence","Notes"
    ], "Inventory and classification of historical sources and production evidence.", {
        "Source Type": ["Mapping","Payload","Salesforce Example","Workfront Example","Intake Example","Validated Test Case","Reference","Other"],
        "Workfront Relevance": ["Yes","No","Unknown"],
        "Mapping Completeness": ["Full","Partial","Unknown"],
        "Source Status": ["Active Evidence","Legacy","Duplicate","Unknown"],
        "Confidence": ["Strong","Moderate","Weak"]
    }),
    ("Mapping Rules", [
        "Rule ID","Status","Mapping Family","Domain","Object / Area","Target Field","Rule Type",
        "Source","Source Field / Path","Population Owner","Transformation / Default","Condition",
        "Required?","Evidence References","Master Version Introduced","Supersedes Rule ID","Notes"
    ], "Reusable approved mapping rules supporting cross-family reasoning and impact analysis.", {
        "Status": ["Candidate","Approved","Superseded"],
        "Domain": ["Salesforce","Workfront"],
        "Rule Type": ["Source","MuleSoft","Salesforce","Default","Transformation","Condition","Cardinality","Other"],
        "Population Owner": ["Source","MuleSoft","Salesforce","Mixed","Unknown"],
        "Required?": ["Yes","No","Conditional","Unknown"]
    }),
    ("Exceptions", [
        "Exception ID","Status","Mapping Family","Partner / Reseller / Context","Domain",
        "Object / Area","Field / Rule","Exception Type","Description","Condition",
        "Reason Not Canonical","Evidence References","Notes"
    ], "Implementation-specific behavior that must remain separate from reusable canonical rules.", {
        "Status": ["Active","Historical","Superseded"],
        "Domain": ["Salesforce","Workfront"],
        "Exception Type": ["Partner","Reseller","One-off","Temporary","Legacy","Other"]
    }),
    ("Clarifications", [
        "Clarification ID","Status","Mapping Family","Domain","Object / Field","Question",
        "Why It Matters","Conflicting / Missing Evidence","Recommended SME","Resolution",
        "Resolution Evidence","Resolved By","Resolved Date","Resulting Decision ID","Resulting Rule ID"
    ], "Open questions and their governed resolution so the same uncertainty is not rediscovered.", {
        "Status": ["Open","Answered","Closed","Superseded"],
        "Domain": ["Salesforce","Workfront","Taxonomy","Other"],
        "Recommended SME": ["Business SME","Salesforce SME","MuleSoft SME","Workfront SME","Solution Architect","Other"]
    }),
    ("Learning & Errors", [
        "Learning ID","Date","Mapping Family","Learning Type","Area","Previous Interpretation",
        "Corrected Interpretation","Why It Was Wrong","Evidence References","Impact",
        "Registry Update Required","Skill / Agent Core Update Required","Status","Validated By","Notes"
    ], "Validated learning from corrections, mistakes, and process improvements.", {
        "Learning Type": ["Domain","Agent Error","Skill Improvement","Governance","Other"],
        "Status": ["Candidate","Validated","Implemented"]
    }),
    ("Versions", [
        "Artifact ID","Canonical Mapping","Version","Status","Created Date","Created By",
        "Approved Date","Approved By","Supersedes Version","File Location",
        "Salesforce Mapping Present","Workfront Mapping Present","Analysis & Assumptions Present",
        "Change Log Present","Change Summary","Analysis Review Reference","Notes"
    ], "Governed Master Mapping version history. Approved versions are immutable.", {
        "Status": ["DRAFT","APPROVED","SUPERSEDED","RETIRED"],
        "Salesforce Mapping Present": ["Yes","No"],
        "Workfront Mapping Present": ["Yes","No"],
        "Analysis & Assumptions Present": ["Yes","No"],
        "Change Log Present": ["Yes","No"]
    }),
]

wb = Workbook()
wb.remove(wb.active)

title_fill = PatternFill("solid", fgColor="1F4E78")
desc_fill = PatternFill("solid", fgColor="D9EAF7")
header_fill = PatternFill("solid", fgColor="5B9BD5")
white_bold = Font(color="FFFFFF", bold=True)
title_font = Font(color="FFFFFF", bold=True, size=14)

for name, headers, desc, validations in SPECS:
    ws = wb.create_sheet(name)
    last_col = get_column_letter(len(headers))
    ws.merge_cells(f"A1:{last_col}1")
    ws["A1"] = f"Onboarding Mapping Agent — {name} Registry"
    ws["A1"].fill = title_fill
    ws["A1"].font = title_font
    ws["A1"].alignment = Alignment(vertical="center")

    ws.merge_cells(f"A2:{last_col}2")
    ws["A2"] = desc
    ws["A2"].fill = desc_fill
    ws["A2"].font = Font(italic=True, color="1F1F1F")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="center")

    for col, header in enumerate(headers, 1):
        cell = ws.cell(4, col, header)
        cell.fill = header_fill
        cell.font = white_bold
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:{last_col}1000"

    for i, header in enumerate(headers, 1):
        h = header.lower()
        width = 18
        if any(k in h for k in ["definition","characteristics","rationale","decision","description","question","interpretation","notes","impact","evidence","context","resolution","change summary","transformation"]):
            width = 28
        elif any(k in h for k in ["date","version","status","domain","confidence","required"]):
            width = 14
        elif any(k in h for k in ["field","object","mapping family","source name","canonical mapping","location"]):
            width = 22
        ws.column_dimensions[get_column_letter(i)].width = width

    for header, values in validations.items():
        if header not in headers:
            continue
        idx = headers.index(header) + 1
        col = get_column_letter(idx)
        formula = '"' + ",".join(values) + '"'
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}5:{col}1000")

    for i, header in enumerate(headers, 1):
        if "Date" in header:
            for row in range(5, 1001):
                ws.cell(row, i).number_format = "yyyy-mm-dd"

    for row in ws.iter_rows(min_row=1, max_row=1000, min_col=1, max_col=len(headers)):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

wb.save(OUT)
print(f"Generated {OUT}")
